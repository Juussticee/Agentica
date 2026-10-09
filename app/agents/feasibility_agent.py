"""Feasibility agent (Phase 2) — spec Section 3 contract.

Consumes the research_agent and books_agent outputs from state (the Supervisor
sets sub_routes = ["research", "books"] first, so both prior answers are
present by the time this node runs), produces a structured study with
Market / Financial / Risk sections, and saves it as a .docx via python-docx.
The .docx path is returned in state["answer"].

Graceful degradation: when books_agent could not contribute (explicit
not-found refusal, missing output, or low-similarity refusal), the study is
still produced — web-only — with a clearly flagged limitation note in both
the .docx and the returned answer, instead of a thin report that pretends
the library was consulted.
"""

import re
import time
from datetime import datetime
from pathlib import Path

from docx import Document

from app.config import NOT_FOUND_PHRASE, REPORTS_DIR
from app.llm_client import get_llm, llm_text
from app.state import AgentState, with_agent_output

_SECTIONS = ("Market", "Financial", "Risk")

_NO_BOOKS_PLACEHOLDER = "(no books output available)"

_PROMPT_TEMPLATE = """You are a business analyst writing a short feasibility study.

Question:
{question}

Findings from web research:
{research_answer}

Findings from the book/company library:
{books_answer}

Write a concise feasibility study with EXACTLY these three sections, in this
order, using ONLY the findings above. Where the findings lack data, state the
gap explicitly instead of inventing numbers. Use markdown headings
"## Market", "## Financial", "## Risk".
"""

_PROMPT_WEB_ONLY_TEMPLATE = """You are a business analyst writing a short feasibility study.

Question:
{question}

Findings from web research:
{research_answer}

The internal book/company library had NO relevant material for this question,
so this study must rest on web findings alone.

Write a concise feasibility study with EXACTLY these three sections, in this
order, using ONLY the findings above. Where the findings lack data, state the
gap explicitly instead of inventing numbers, and note where internal company
data would normally be required. Use markdown headings
"## Market", "## Financial", "## Risk".
"""


def _books_usable(books_answer: str) -> bool:
    """True when the books agent produced real content we can synthesize against."""
    text = (books_answer or "").strip()
    if not text or text == _NO_BOOKS_PLACEHOLDER:
        return False
    return NOT_FOUND_PHRASE.lower() not in text.lower()


def _extract_section(study: str, name: str) -> str:
    """Pull one section body out of the model's markdown output (best effort)."""
    header_re = re.compile(
        rf"^\s*(?:#+\s*{name}\b|\**{name}\**\s*:?).*$",
        re.IGNORECASE | re.MULTILINE,
    )
    match = header_re.search(study)
    if not match:
        return ""
    start = match.end()
    others = "|".join(_SECTIONS)
    next_re = re.search(
        rf"^\s*(?:#+\s*(?:{others})\b|\**(?:{others})\**\s*:?).*$",
        study[start:], re.IGNORECASE | re.MULTILINE,
    )
    end = start + next_re.start() if next_re else len(study)
    return study[start:end].strip()


def _write_docx(question: str, study: str, limitation: str) -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = REPORTS_DIR / f"feasibility_{stamp}.docx"

    doc = Document()
    doc.add_heading(f"Feasibility Study: {question}", 0)
    doc.add_paragraph(
        f"Generated {datetime.now():%Y-%m-%d %H:%M} by the Master Agent System "
        f"(research + books agents, feasibility synthesis)."
    )
    if limitation:
        doc.add_heading("Limitations", level=1)
        doc.add_paragraph(limitation)
    for section in _SECTIONS:
        doc.add_heading(section, level=1)
        body = _extract_section(study, section)
        doc.add_paragraph(body if body else "(section not produced by the model)")
    doc.save(path)
    return path


def feasibility_agent(state: AgentState) -> dict:
    question = state.get("question", "")
    answers = state.get("agent_answers") or {}
    research_answer = answers.get("research") or "(no research output available)"
    books_answer = answers.get("books") or _NO_BOOKS_PLACEHOLDER

    web_only = not _books_usable(books_answer)
    template = _PROMPT_WEB_ONLY_TEMPLATE if web_only else _PROMPT_TEMPLATE
    limitation = (
        "Web-only study: the internal book/company library had no relevant "
        "material for this question, so findings are not corroborated against "
        "internal sources." if web_only else ""
    )

    start = time.perf_counter()
    study = llm_text(
        get_llm().invoke(
            template.format(
                question=question,
                research_answer=research_answer,
                books_answer=books_answer,
            )
        )
    ).strip()
    elapsed = time.perf_counter() - start

    path = _write_docx(question, study, limitation)
    mode = "web-only" if web_only else "research+books"
    print(f"[feasibility_agent] {mode} study written in {elapsed:.1f}s -> {path.name}")

    answer = f"Feasibility study saved to: {path}\n\n{study}"
    if limitation:
        answer = f"⚠ {limitation}\n\n{answer}"
    return with_agent_output(state, "feasibility", answer, ["feasibility"])
